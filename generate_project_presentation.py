from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR


WIDTH = Inches(13.333)
HEIGHT = Inches(7.5)
NAVY = RGBColor(10, 25, 40)
NAVY_2 = RGBColor(20, 39, 59)
BLUE = RGBColor(67, 136, 255)
CYAN = RGBColor(73, 214, 255)
TEAL = RGBColor(22, 189, 202)
WHITE = RGBColor(248, 250, 252)
GRAY = RGBColor(148, 163, 184)
LIGHT = RGBColor(226, 232, 240)
GREEN = RGBColor(34, 197, 94)
AMBER = RGBColor(245, 158, 11)


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, title, subtitle=None, accent=True):
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(8.8), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.name = 'Aptos'
    run.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT

    if accent:
        accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.25), Inches(1.0), Inches(0.08))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = CYAN
        accent_bar.line.fill.background()

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.75), Inches(1.5), Inches(10.5), Inches(0.6))
        sub_tf = sub.text_frame
        sub_tf.word_wrap = True
        p2 = sub_tf.paragraphs[0]
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.size = Pt(11)
        run2.font.color.rgb = LIGHT
        run2.font.name = 'Aptos'
        p2.alignment = PP_ALIGN.LEFT


def add_metric_card(slide, x, y, w, h, title, value, accent_color, subtitle=''):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = NAVY_2
    box.line.color.rgb = accent_color
    box.line.width = Pt(1.2)

    title_box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.15), w - Inches(0.36), Inches(0.35))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(10)
    r.font.bold = False
    r.font.color.rgb = LIGHT
    r.font.name = 'Aptos'

    val_box = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.42), w - Inches(0.36), Inches(0.6))
    tf2 = val_box.text_frame
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    r2.text = value
    r2.font.size = Pt(22)
    r2.font.bold = True
    r2.font.color.rgb = WHITE
    r2.font.name = 'Aptos'

    if subtitle:
        sbox = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.92), w - Inches(0.36), Inches(0.25))
        stf = sbox.text_frame
        sp = stf.paragraphs[0]
        sr = sp.add_run()
        sr.text = subtitle
        sr.font.size = Pt(8)
        sr.font.color.rgb = GRAY
        sr.font.name = 'Aptos'


def add_card(slide, x, y, w, h, title, body, accent_color=CYAN):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = NAVY_2
    box.line.color.rgb = accent_color
    box.line.width = Pt(1.2)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent_color
    bar.line.fill.background()

    t = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.18), w - Inches(0.36), Inches(0.35))
    tf = t.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = 'Aptos'

    b = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.55), w - Inches(0.36), h - Inches(0.65))
    btf = b.text_frame
    btf.word_wrap = True
    p2 = btf.paragraphs[0]
    p2.text = body
    p2.font.size = Pt(11)
    p2.font.name = 'Aptos'
    p2.font.color.rgb = LIGHT


def add_process_box(slide, x, y, w, h, label, color=BLUE):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.color.rgb = color
    box.line.width = Pt(1)

    t = slide.shapes.add_textbox(x + Inches(0.1), y + Inches(0.18), w - Inches(0.2), h - Inches(0.2))
    tf = t.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = 'Aptos'


def add_arrow(slide, x1, y1, x2, y2):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x1, y1, x2 - x1, y2 - y1)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = CYAN
    arrow.line.fill.background()


prs = Presentation()
prs.slide_width = WIDTH
prs.slide_height = HEIGHT

# Slide 1 - Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)

# top-right accent circle
circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.8), Inches(0.55), Inches(2.2), Inches(2.2))
circle.fill.solid(); circle.fill.fore_color.rgb = BLUE; circle.line.fill.background()

# left text
shape = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(7.4), Inches(1.2))
tf = shape.text_frame
p = tf.paragraphs[0]
r = p.add_run()
r.text = 'JOB MARKET'
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = CYAN
r.font.name = 'Aptos'

shape2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(8.5), Inches(1.8))
tf2 = shape2.text_frame
p2 = tf2.paragraphs[0]
r2 = p2.add_run()
r2.text = 'ANALYTICS PLATFORM'
r2.font.size = Pt(28)
r2.font.bold = True
r2.font.color.rgb = WHITE
r2.font.name = 'Aptos'

sub = slide.shapes.add_textbox(Inches(0.8), Inches(3.1), Inches(7.2), Inches(1.0))
sub_tf = sub.text_frame
p3 = sub_tf.paragraphs[0]
r3 = p3.add_run()
r3.text = 'Data-driven insights into hiring trends, skill demand, and labor market movement.'
r3.font.size = Pt(16)
r3.font.color.rgb = LIGHT
r3.font.name = 'Aptos'

# metric bar
box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.5), Inches(3.2), Inches(0.8))
box.fill.solid(); box.fill.fore_color.rgb = BLUE; box.line.fill.background()
box_text = slide.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(2.8), Inches(0.5))
box_tf = box_text.text_frame
p4 = box_tf.paragraphs[0]
r4 = p4.add_run(); r4.text = 'ETL + AI + Insights'; r4.font.size = Pt(15); r4.font.bold = True; r4.font.color.rgb = WHITE

# side panel on right
panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.2), Inches(1.5), Inches(3.4), Inches(3.8))
panel.fill.solid(); panel.fill.fore_color.rgb = NAVY_2; panel.line.color.rgb = CYAN; panel.line.width = Pt(1.2)

mini1 = slide.shapes.add_textbox(Inches(9.6), Inches(2.0), Inches(2.4), Inches(0.5))
mini1_tf = mini1.text_frame
p5 = mini1_tf.paragraphs[0]
r5 = p5.add_run(); r5.text = 'Skills'; r5.font.size = Pt(14); r5.font.bold = True; r5.font.color.rgb = CYAN

mini2 = slide.shapes.add_textbox(Inches(9.6), Inches(2.7), Inches(2.4), Inches(0.5))
mini2_tf = mini2.text_frame
p6 = mini2_tf.paragraphs[0]
r6 = p6.add_run(); r6.text = 'Locations'; r6.font.size = Pt(14); r6.font.bold = True; r6.font.color.rgb = WHITE

mini3 = slide.shapes.add_textbox(Inches(9.6), Inches(3.4), Inches(2.4), Inches(0.5))
mini3_tf = mini3.text_frame
p7 = mini3_tf.paragraphs[0]
r7 = p7.add_run(); r7.text = 'Hiring Trends'; r7.font.size = Pt(14); r7.font.bold = True; r7.font.color.rgb = LIGHT

mini4 = slide.shapes.add_textbox(Inches(9.6), Inches(4.1), Inches(2.4), Inches(0.5))
mini4_tf = mini4.text_frame
p8 = mini4_tf.paragraphs[0]
r8 = p8.add_run(); r8.text = 'AI Recommendations'; r8.font.size = Pt(12); r8.font.bold = True; r8.font.color.rgb = CYAN

# Slide 2 - Overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'Project Overview', 'Turning raw job data into actionable market intelligence.')

add_card(slide, Inches(0.8), Inches(2.0), Inches(3.5), Inches(2.1), 'Problem', 'The job market changes fast. Employers, candidates, and analysts need more than static listings—they need trends, signals, and decisions backed by data.', CYAN)
add_card(slide, Inches(4.7), Inches(2.0), Inches(3.8), Inches(2.1), 'Solution', 'This platform collects, cleans, and analyzes job market data to create a real-time, searchable view of hiring demand and skill evolution.', TEAL)
add_card(slide, Inches(8.9), Inches(2.0), Inches(3.6), Inches(2.1), 'Value', 'It supports smarter recruiting, more focused career planning, and deeper understanding of where industries and skills are moving.', BLUE)

add_metric_card(slide, Inches(1.0), Inches(5.0), Inches(2.8), Inches(1.2), 'Data Sources', 'Multiple', CYAN, 'APIs, jobs, market feeds')
add_metric_card(slide, Inches(4.2), Inches(5.0), Inches(2.8), Inches(1.2), 'Processing', 'Clean + normalize', BLUE, 'Structured for analysis')
add_metric_card(slide, Inches(7.4), Inches(5.0), Inches(2.8), Inches(1.2), 'Insights', 'AI guided', TEAL, 'Skills, trends, patterns')

# Slide 3 - Objectives
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'Objectives', 'Building a smart ecosystem around labor market intelligence.')

objectives = [
    'Collect job-related data from multiple sources and maintain consistent updates.',
    'Transform raw information into structured, query-friendly datasets.',
    'Identify skill trends, hiring hotspots, and role demand patterns.',
    'Support career and recruitment decisions with actionable analytics.',
    'Provide a dashboard experience that makes complex signals easy to understand.',
    'Integrate AI features for recommendations and conversational querying.'
]

for i, obj in enumerate(objectives):
    y = Inches(1.9 + (i % 3) * 1.3)
    x = Inches(0.8 + (i // 3) * 6.1)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.4), Inches(1.0))
    card.fill.solid(); card.fill.fore_color.rgb = NAVY_2; card.line.color.rgb = CYAN; card.line.width = Pt(1.2)
    text = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.12), Inches(5.0), Inches(0.7))
    tf = text.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = f'{i+1}. {obj}'; r.font.size = Pt(12); r.font.color.rgb = WHITE; r.font.name = 'Aptos' 

# Slide 4 - Architecture
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'System Architecture', 'A connected workflow from raw market signals to intelligent decision-making.')

add_process_box(slide, Inches(0.8), Inches(2.15), Inches(1.8), Inches(1.0), 'Data Sources', BLUE)
add_process_box(slide, Inches(2.9), Inches(2.15), Inches(1.9), Inches(1.0), 'ETL Pipeline', TEAL)
add_process_box(slide, Inches(5.2), Inches(2.15), Inches(1.9), Inches(1.0), 'Database', GREEN)
add_process_box(slide, Inches(7.6), Inches(2.15), Inches(2.1), Inches(1.0), 'AI & Analytics', CYAN)
add_process_box(slide, Inches(10.1), Inches(2.15), Inches(2.2), Inches(1.0), 'Dashboard', BLUE)

# connectors arrows
add_arrow(slide, Inches(2.6), Inches(2.65), Inches(2.9), Inches(2.65))
add_arrow(slide, Inches(4.8), Inches(2.65), Inches(5.2), Inches(2.65))
add_arrow(slide, Inches(7.1), Inches(2.65), Inches(7.6), Inches(2.65))
add_arrow(slide, Inches(9.7), Inches(2.65), Inches(10.1), Inches(2.65))

# lower layer cards
add_card(slide, Inches(0.9), Inches(4.1), Inches(3.6), Inches(1.8), 'Input Layer', 'Job boards, public APIs, and market data feeds bring real-time hiring information into the system.', CYAN)
add_card(slide, Inches(4.8), Inches(4.1), Inches(3.6), Inches(1.8), 'Processing Layer', 'Extraction, cleaning, normalization, and transformation standardize structured insights for analysis.', TEAL)
add_card(slide, Inches(8.7), Inches(4.1), Inches(3.6), Inches(1.8), 'Output Layer', 'Analytics, recommendations, and visual dashboards turn data into business-ready decisions.', BLUE)

# Slide 5 - Core Features
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'Core Features', 'The project blends analytics, intelligence, and usability in one ecosystem.')

features = [
    ('Real-time collection', 'Scheduled and ongoing ingestion keeps the platform updated with current labor market data.'),
    ('Skill taxonomy', 'Classification and mapping of roles and skill sets make market demand easier to interpret.'),
    ('Trend analytics', 'Filters for industry, location, and role reveal changing patterns across the market.'),
    ('AI-driven insights', 'Natural-language queries and recommendations help users explore complex hiring trends faster.'),
    ('Dashboard interface', 'Web-based views convert data into effective charts, summaries, and decision support.'),
    ('Scalable backend', 'A structured architecture supports future growth, new data sources, and richer analytics.'),
]

for i, (title, body) in enumerate(features):
    x = Inches(0.75 + (i % 3) * 4.0)
    y = Inches(1.9 + (i // 3) * 2.1)
    add_card(slide, x, y, Inches(3.5), Inches(1.6), title, body, CYAN if i % 2 == 0 else TEAL)

# Slide 6 - Data Flow
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'Data Flow Design', 'A clean pipeline that turns raw signals into structured intelligence.')

add_process_box(slide, Inches(0.7), Inches(2.2), Inches(2.0), Inches(1.1), '1. Collect', BLUE)
add_process_box(slide, Inches(2.9), Inches(2.2), Inches(2.2), Inches(1.1), '2. Clean', TEAL)
add_process_box(slide, Inches(5.4), Inches(2.2), Inches(2.1), Inches(1.1), '3. Store', GREEN)
add_process_box(slide, Inches(8.0), Inches(2.2), Inches(2.0), Inches(1.1), '4. Analyze', CYAN)
add_process_box(slide, Inches(10.5), Inches(2.2), Inches(2.1), Inches(1.1), '5. Visualize', BLUE)

for i in range(4):
    add_arrow(slide, Inches(2.7 + i * 2.5), Inches(2.75), Inches(2.9 + i * 2.5), Inches(2.75))

add_card(slide, Inches(1.0), Inches(4.2), Inches(5.0), Inches(1.7), 'Pipeline Logic', 'Collection is consistent and repeatable, transformation removes noise and standardizes values, and storage ensures clean, queryable datasets for deeper analysis.', CYAN)
add_card(slide, Inches(6.9), Inches(4.2), Inches(5.2), Inches(1.7), 'Analytical Outcome', 'This structure becomes the foundation for dashboards, role comparisons, trend forecasting, and AI-supported recommendations.', TEAL)

# Slide 7 - AI Layer
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'AI Intelligence Layer', 'Giving users a smarter way to explore the market.')

add_card(slide, Inches(0.9), Inches(1.9), Inches(3.6), Inches(2.0), 'Skill Mapping', 'The system identifies key skills, relationships between roles, and skills that are rising in demand across industries.', CYAN)
add_card(slide, Inches(4.9), Inches(1.9), Inches(3.6), Inches(2.0), 'Recommendation Engine', 'AI can highlight market opportunities and suggest roles or skill development paths based on trends.', TEAL)
add_card(slide, Inches(8.9), Inches(1.9), Inches(3.5), Inches(2.0), 'Chatbot Queries', 'Users can ask natural-language questions to understand demand, gaps, or hiring patterns quickly.', BLUE)

add_metric_card(slide, Inches(1.3), Inches(4.7), Inches(2.8), Inches(1.3), 'Trend Signals', 'Market Pulse', CYAN, 'Role + skill demand')
add_metric_card(slide, Inches(4.8), Inches(4.7), Inches(2.8), Inches(1.3), 'Opportunity Fit', 'Match Score', TEAL, 'Candidate vs role fit')
add_metric_card(slide, Inches(8.3), Inches(4.7), Inches(2.8), Inches(1.3), 'Insight Depth', 'Conversational', BLUE, 'AI explanations')

# Slide 8 - Dashboard Experience
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'Dashboard Experience', 'A user-facing layer built for clarity and fast decisions.')

# mock dashboard cards
panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.7), Inches(11.9), Inches(4.9))
panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(15, 32, 49); panel.line.color.rgb = CYAN; panel.line.width = Pt(1)

add_metric_card(slide, Inches(1.0), Inches(2.1), Inches(2.4), Inches(1.3), 'Open Roles', '12.4k', CYAN, '+8.2% MoM')
add_metric_card(slide, Inches(3.8), Inches(2.1), Inches(2.4), Inches(1.3), 'Skill Demand', 'AI/ML', TEAL, 'High growth')
add_metric_card(slide, Inches(6.6), Inches(2.1), Inches(2.4), Inches(1.3), 'Regions', 'Top 5', BLUE, 'Hiring hubs')
add_metric_card(slide, Inches(9.4), Inches(2.1), Inches(2.4), Inches(1.3), 'Forecast', 'Positive', GREEN, 'Next quarter')

# simplified chart-like blocks
for i in range(6):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2 + i*1.7), Inches(4.8), Inches(0.8), Inches(0.9 + (i % 5) * 0.25))
    bar.fill.solid(); bar.fill.fore_color.rgb = [CYAN, BLUE, TEAL, GREEN, AMBER, CYAN][i]; bar.line.fill.background()

add_card(slide, Inches(9.4), Inches(4.3), Inches(2.7), Inches(1.2), 'Dashboard Goal', 'Make market signals actionable, not overwhelming.', CYAN)

# Slide 9 - Unique Differentiators
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'What Makes It Unique', 'Designing the project as a market intelligence platform, not just a data extractor.')

unique = [
    ('Market-first lens', 'It focuses on labor movement, hiring demand, and skill evolution rather than isolated job postings.'),
    ('AI-assisted exploration', 'Users can query trends in plain language and discover patterns quickly.'),
    ('Modular architecture', 'The system is structured for future expansion, additional sources, and stronger analytics.'),
    ('Decision-ready output', 'The dashboard is designed for recruiters, job seekers, and analysts to act on insights without technical friction.')
]
for i, (title, body) in enumerate(unique):
    x = Inches(0.8 + (i % 2) * 6.1)
    y = Inches(1.9 + (i // 2) * 2.2)
    add_card(slide, x, y, Inches(5.3), Inches(1.8), title, body, CYAN if i % 2 == 0 else BLUE)

# Slide 10 - Impact and Next Steps
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, NAVY)
add_title(slide, 'Impact & Next Steps', 'A platform ready to support better decisions in a changing job market.')

add_card(slide, Inches(0.95), Inches(2.0), Inches(5.3), Inches(2.5), 'Business Impact', 'This platform helps employers understand talent demand, supports job seekers with realistic market signals, and gives analysts a clearer view of market structure.', CYAN)
add_card(slide, Inches(6.8), Inches(2.0), Inches(5.3), Inches(2.5), 'Next Phase', 'Expand data sources, refine AI interpretation, add forecasting, and improve the dashboard with deeper role and region segmentation.', TEAL)

add_metric_card(slide, Inches(1.3), Inches(5.1), Inches(3.0), Inches(1.2), 'Use Cases', 'Recruiting', CYAN, 'Hiring trend analysis')
add_metric_card(slide, Inches(4.8), Inches(5.1), Inches(3.0), Inches(1.2), 'Use Cases', 'Career Planning', TEAL, 'Skill pathways')
add_metric_card(slide, Inches(8.3), Inches(5.1), Inches(3.0), Inches(1.2), 'Use Cases', 'Market Research', BLUE, 'Industry forecasting')

prs.core_properties.title = 'Job Market Analytics Project'
prs.core_properties.subject = 'Project overview and architecture presentation'
prs.core_properties.author = 'GitHub Copilot'
prs.core_properties.company = 'Job Market Analytics'

output_file = r'c:\Users\DELL\Job-Market-Analytics\docs\Job_Market_Analytics_Presentation.pptx'
prs.save(output_file)
print(f'Presentation created: {output_file}')
